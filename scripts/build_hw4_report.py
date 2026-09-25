"""Build the editable Homework 4 report from verified code, data, and screenshots."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
HW4 = ROOT / "reports" / "hw04"
SHOTS = HW4 / "screenshots"
OUTPUT = HW4 / "Vartak_HW4_report.docx"


def lines(path: str, start: int, end: int) -> str:
    source = (ROOT / path).read_text(encoding="utf-8").splitlines()
    return "\n".join(source[start - 1:end])


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_code(doc: Document, label: str, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    lead = p.add_run(label + "\n")
    lead.bold = True
    lead.font.name = "Arial"
    lead.font.size = Pt(8)
    run = p.add_run(code)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(7.2)
    p.paragraph_format.line_spacing = 0.9
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F4F6")
    p_pr.append(shading)


def add_figure(doc: Document, filename: str, number: int, caption: str, width=6.45) -> None:
    image_path = SHOTS / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    run = cap.add_run(f"Figure {number}. {caption}")
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(8.5)


def add_unit(doc: Document, heading: str, explanation: str, code_label: str, code: str,
             filename: str, figure: int, caption: str, break_before: bool = True) -> None:
    if break_before:
        doc.add_page_break()
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph(explanation)
    p.paragraph_format.space_after = Pt(4)
    add_code(doc, code_label, code)
    add_figure(doc, filename, figure, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[index])
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if widths:
                cells[index].width = Inches(widths[index])
            if row_index % 2:
                shade(cells[index], "EEF4F8")
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8.5)
            cells[index].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        page_number(section.footer.paragraphs[0])


def build(commit_hash: str) -> None:
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DATA 260 Homework 4")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("React, MySQL, Query Tuning, and Grounded RAG")
    subtitle_run.bold = True
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(15)
    doc.add_paragraph()
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Vedant Vartak\nGrocery Supply and Recall Notices").bold = True

    doc.add_heading("Personal configuration", level=1)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["SID4", "3539"],
            ["PORT_BASE", "8839"],
            ["PREFIX", "s3539"],
            ["SEED", "3539"],
            ["VERIFY_SEED", "263539"],
            ["DOMAIN_ID", "3"],
            ["Assigned domain", "Grocery Supply and Recall Notices"],
            ["Hardware", "Apple M4 MacBook Air, 10 CPU cores, 16 GB memory"],
            ["Local LLM", "qwen3:8b"],
            ["Embedding model", "sentence-transformers/all-MiniLM-L6-v2"],
            ["Code and results commit", commit_hash],
            ["Final submission tag", "hw4"],
        ],
        widths=[2.0, 4.5],
    )
    p = doc.add_paragraph()
    p.add_run("GitHub repository: ").bold = True
    p.add_run("https://github.com/vedantvartakwork/data260-3539")
    doc.add_paragraph(
        "The repository link was checked, and both required collaborators, Sbnikitha and "
        "supriyaselvanganesan, have access. The Canvas copy is named Vartak_HW4.pdf."
    )
    doc.add_paragraph(
        "This submission extends the existing course repository with a protected React client, "
        "FastAPI and MySQL persistence, measured N+1 query behavior, an index experiment, and a "
        "grounded local RAG evaluation. The final smoke test passed all 11 objective checks."
    )

    doc.add_page_break()
    doc.add_heading("Part 1 React client", level=1)
    doc.add_paragraph(
        "The client uses React Router, hooks, props, and protected routes. Login state is recovered "
        "from the backend session, and every create, update, and delete operation is persisted through "
        "the API before the user returns to the home page."
    )

    add_unit(doc, "1.1 Login and protected routing",
             "The login component stores the email and password with useState and passes them to the parent handler. ProtectedRoute redirects unauthenticated users to /login.",
             "Login.jsx and App.jsx", lines("code/web_application/frontend/src/components/Login.jsx", 4, 21) + "\n\n" + lines("code/web_application/frontend/src/App.jsx", 12, 17),
             "01_login_page.png", 1, "Login page with email and password fields.", break_before=False)
    add_unit(doc, "1.2 Home record list",
             "Home receives records and the authenticated user through props. It renders the current recall records and routes each row to update and delete actions.",
             "Home.jsx", lines("code/web_application/frontend/src/components/Home.jsx", 52, 72),
             "02_home_record_list.png", 2, "Protected home page showing recall records and related-event counts.")
    add_unit(doc, "1.3 Create record form",
             "CreateRecord receives the create callback as a prop and reuses the validated RecordForm component.",
             "CreateRecord.jsx", lines("code/web_application/frontend/src/components/CreateRecord.jsx", 5, 14),
             "03_create_form.png", 3, "Create-record route and domain-specific form.")
    add_unit(doc, "1.4 Create operation result",
             "After the POST succeeds, the client reloads the list, sets a success message with the new auto-incremented ID, and redirects to the root route.",
             "App.jsx", lines("code/web_application/frontend/src/App.jsx", 65, 70),
             "04_create_success.png", 4, "Successful creation of recall ID 5003.")
    add_unit(doc, "1.5 Update record form",
             "The update component reads the route ID, loads the current record in useEffect, and passes the edited payload back through the onUpdate prop.",
             "UpdateRecord.jsx", lines("code/web_application/frontend/src/components/UpdateRecord.jsx", 8, 30),
             "05_update_form.png", 5, "Update form populated from MySQL for recall ID 5003.")
    add_unit(doc, "1.6 Update operation result",
             "The update handler writes the record through the backend, refreshes the list, and redirects to the home route with confirmation.",
             "App.jsx", lines("code/web_application/frontend/src/App.jsx", 72, 77),
             "06_update_success.png", 6, "Successful update confirmation on the home page.")
    add_unit(doc, "1.7 Delete confirmation",
             "The delete route loads the selected record and requires a deliberate button click before calling the parent delete handler.",
             "DeleteRecord.jsx", lines("code/web_application/frontend/src/components/DeleteRecord.jsx", 7, 30),
             "07_delete_confirmation.png", 7, "Delete confirmation for recall ID 5003.")
    add_unit(doc, "1.8 Delete operation result",
             "After the DELETE request succeeds, the client reloads the list and shows the deleted ID in a success message.",
             "App.jsx", lines("code/web_application/frontend/src/App.jsx", 79, 84),
             "08_delete_success.png", 8, "Successful deletion confirmation on the home page.")

    doc.add_page_break()
    doc.add_heading("Part 2 MySQL persistence and server-side sessions", level=1)
    doc.add_paragraph(
        "The FastAPI service uses SQLAlchemy with the required db_session_basede26 variable and the "
        "s3539_rel MySQL database. Passwords are stored as hashes. Login creates a random opaque token "
        "in the sessions table and sends only that token in an HTTP-only SameSite cookie. Every recall "
        "endpoint depends on require_user."
    )
    add_unit(doc, "2.1 Database connection and required variable",
             "The connection targets s3539_rel and exposes the exact required SQLAlchemy session variable name.",
             "db.py", lines("code/web_application/db.py", 12, 36),
             "21_mysql_database_evidence.png", 9, "MySQL database, table, row-count, user, and session evidence.", break_before=False)
    add_unit(doc, "2.2 Project folder structure",
             "All application code remains in shared root-level folders and extends the existing repository rather than creating a separate homework copy.",
             "Command used to generate the evidence", "find code src sql scripts tests reports/hw04 -maxdepth 4 -type f | sort",
             "22_project_folder_structure.png", 10, "Project structure covering backend, React, SQL, scripts, tests, and report artifacts.")
    add_unit(doc, "2.3 Login API and opaque session cookie",
             "Login verifies the password hash, replaces any earlier session, stores a random token server-side, and marks the cookie HTTP-only.",
             "api_auth.py", lines("code/web_application/routers/api_auth.py", 50, 84),
             "09_postman_login.png", 11, "Postman login response with authenticated user data.")
    add_unit(doc, "2.4 View all records",
             "The protected list endpoint eager-loads related events and supports a bounded page size and optional product or brand search.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 104, 137),
             "10_postman_list.png", 12, "GET recall list response with fixed implementation and one SQL statement.")
    add_unit(doc, "2.5 Create a record",
             "The POST endpoint validates the payload, adds a RecallNotice, commits it, refreshes its auto-incremented ID, and returns HTTP 201.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 156, 166),
             "11_postman_create.png", 13, "POST response for newly created recall ID 5004.")
    add_unit(doc, "2.6 View a record by ID",
             "The ID endpoint eager-loads events and returns 404 if the record does not exist.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 140, 153),
             "12_postman_read.png", 14, "GET response for the created recall.")
    add_unit(doc, "2.7 Update a record",
             "The PUT endpoint loads the row, applies validated fields, commits, and returns the current related-event list.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 169, 190),
             "13_postman_update.png", 15, "PUT response showing the updated brand and recall details.")
    add_unit(doc, "2.8 Delete a record",
             "The DELETE endpoint removes the requested row, returns 404 when no row was affected, and otherwise returns HTTP 204.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 193, 205),
             "14_postman_delete.png", 16, "DELETE response with HTTP 204 No Content.")

    doc.add_page_break()
    doc.add_heading("Part 3 N plus 1 measurement and query tuning", level=1)
    doc.add_paragraph(
        "The deterministic seed script creates 5,000 recall notices and 200 related recall events with "
        "SEED 3539. The naive endpoint runs one main query plus one event query per returned row. The "
        "fixed endpoint eager-loads the relationship in one SQL statement. Each implementation was "
        "measured 30 times at page sizes 10, 50, and 200, giving 180 saved raw requests."
    )
    add_code(doc, "seed_hw04.py", lines("scripts/seed_hw04.py", 17, 26) + "\n...\n" + lines("scripts/seed_hw04.py", 75, 93))
    doc.add_paragraph("The database screenshot in Figure 9 verifies the seeded counts and authentication tables.")

    add_unit(doc, "3.1 Naive endpoint at page size 10",
             "For each returned recall, the loop sends a separate SELECT for recall_events. Ten rows therefore require 11 SQL statements.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 43, 72),
             "15_naive_10.png", 17, "Naive endpoint at page size 10 with 11 SQL queries.")
    add_unit(doc, "3.2 Fixed endpoint at page size 10",
             "joinedload fetches the same relationship eagerly, so the query count remains one.",
             "recalls.py", lines("code/web_application/routers/recalls.py", 78, 101),
             "16_fixed_10.png", 18, "Fixed eager-loading endpoint at page size 10 with one SQL query.")
    add_unit(doc, "3.3 Naive endpoint at page size 50",
             "The N+1 pattern grows with the page: one recall query plus 50 related-event queries.",
             "Measured request", "GET /api/v1/recalls/naive?page_size=50\nExpected SQL statements: 1 + page_size = 51",
             "17_naive_50.png", 19, "Naive endpoint at page size 50 with 51 SQL queries.")
    add_unit(doc, "3.4 Fixed endpoint at page size 50",
             "The eager-loading implementation returns the same record and event structure using one SQL statement.",
             "Measured request", "GET /api/v1/recalls/fixed?page_size=50\nExpected SQL statements: 1",
             "18_fixed_50.png", 20, "Fixed endpoint at page size 50 with one SQL query.")
    add_unit(doc, "3.5 Naive endpoint at page size 200",
             "At the largest page size, the naive endpoint issues 201 statements and shows the linear query-growth problem clearly.",
             "Measured request", "GET /api/v1/recalls/naive?page_size=200\nExpected SQL statements: 1 + page_size = 201",
             "19_naive_200.png", 21, "Naive endpoint at page size 200 with 201 SQL queries.")
    add_unit(doc, "3.6 Fixed endpoint at page size 200",
             "The fixed endpoint still performs one statement even when 200 records are returned.",
             "Measured request", "GET /api/v1/recalls/fixed?page_size=200\nExpected SQL statements: 1",
             "20_fixed_200.png", 22, "Fixed endpoint at page size 200 with one SQL query.")

    doc.add_page_break()
    doc.add_heading("3.7 Latency and SQL statement results", level=2)
    add_code(doc, "Experiment loop", lines("scripts/run_hw4_nplus1_experiment.py", 42, 61))
    add_table(
        doc,
        ["Page size", "Version", "SQL/req", "p50 ms", "p95 ms", "p99 ms", "Speed-up"],
        [
            ["10", "naive", "11", "6.023", "7.935", "8.608", "-"],
            ["10", "fixed", "1", "3.303", "8.523", "25.920", "1.82x"],
            ["50", "naive", "51", "18.035", "20.308", "22.296", "-"],
            ["50", "fixed", "1", "5.656", "7.242", "19.120", "3.19x"],
            ["200", "naive", "201", "59.150", "72.617", "74.108", "-"],
            ["200", "fixed", "1", "11.273", "14.457", "20.258", "5.25x"],
        ],
        widths=[0.7, 0.8, 0.75, 0.8, 0.8, 0.8, 0.8],
    )
    doc.add_paragraph(
        "The fixed median speed-up grows from 1.82x at 10 rows to 5.25x at 200 rows because the naive "
        "version adds one round trip and one query execution for every returned record. The fixed query "
        "count stays constant, so its database overhead grows much more slowly. The p95 and p99 values "
        "still vary because local scheduling, connection reuse, serialization, and warm-up noise affect "
        "individual requests; the median and query counts show the intended scaling behavior most clearly."
    )
    add_figure(doc, "23_nplus1_metrics.png", 23, "Saved N+1 metrics, index result, and RAG summary.")

    add_unit(doc, "3.8 Index and EXPLAIN comparison",
             "The experiment drops any earlier copy of the index, captures EXPLAIN, creates an index on category, and captures EXPLAIN again.",
             "run_hw4_explain.py", lines("scripts/run_hw4_explain.py", 27, 46),
             "24_explain_before_after_index.png", 24, "EXPLAIN before and after adding idx_recall_notices_category.")
    doc.add_paragraph(
        "Before the index, MySQL used access type ALL, no key, and estimated 4,967 rows, which indicates a "
        "full table scan. After the index, the access type changed to ref, MySQL selected "
        "idx_recall_notices_category, and the estimate fell to 1,188 rows. The query can now locate rows "
        "through the category index instead of scanning the full table."
    )

    doc.add_page_break()
    doc.add_heading("Part 4 Grounded RAG question answering", level=1)
    doc.add_paragraph(
        "The RAG system uses five FDA recall documents, 500-token chunks with 50-token overlap, source and "
        "chunk identifiers, all-MiniLM-L6-v2 embeddings, a LlamaIndex SimpleVectorStore, and qwen3:8b. "
        "The same six questions are evaluated with no RAG, basic top-3 RAG, and context-engineered RAG."
    )
    add_unit(doc, "4.1 Corpus, chunks, metadata, and vector index",
             "Every document stores its source name. Each generated node receives a stable chunk_id before it is embedded and loaded into the vector store.",
             "rag.py", lines("rag.py", 38, 49) + "\n\n" + lines("rag.py", 60, 80),
             "25_rag_corpus_and_index.png", 25, "Five-document corpus, 249 chunks, metadata, embedding model, and vector store.", break_before=False)
    add_unit(doc, "4.2 Top-k retrieval output",
             "Retrieval records rank, score, source, chunk ID, text, and latency before the model is called. The printed evidence covers all six question types.",
             "rag.py", lines("rag.py", 83, 99) + "\n\n" + lines("rag.py", 207, 215),
             "26_top_k_retrieval_printout.png", 26, "Top-3 chunks, sources, scores, and IDs printed for all six questions.")
    add_unit(doc, "4.3 No RAG, basic RAG, and context-engineered RAG",
             "The context-engineered prompt labels sources, prohibits outside knowledge, requires citations, and gives the exact refusal sentence.",
             "rag.py", lines("rag.py", 137, 158),
             "27_rag_configurations_and_refusals.png", 27, "Three-configuration comparison and exact refusals for Q5 and Q6.")
    add_unit(doc, "4.4 Six-question evaluation",
             "Evaluation records retrieval, correctness, grounding, refusal behavior, and format compliance for every question and configuration.",
             "rag.py", lines("rag.py", 174, 204),
             "28_rag_six_question_evaluation.png", 28, "Evaluation table with context-RAG totals of 6/6 for correctness, grounding, and format.")

    doc.add_page_break()
    doc.add_heading("4.5 Evaluation summary", level=2)
    add_table(
        doc,
        ["Configuration", "Correct answers", "Grounded answers", "Format compliant"],
        [
            ["No RAG", "1/6", "0/6", "4/6"],
            ["Basic RAG", "4/6", "0/6", "4/6"],
            ["Context RAG", "6/6", "6/6", "6/6"],
        ],
        widths=[1.8, 1.5, 1.5, 1.5],
    )
    doc.add_paragraph(
        "Context RAG was the only configuration that was fully correct, cited the supplied evidence, "
        "and refused both unsupported questions with the required wording."
    )

    add_unit(doc, "4.6 Context-size sweep",
             "The two-chunk question was repeated at k=1, k=3, and k=5. The larger setting supplied the second required public-warning chunk, while also adding a less relevant retail-consignee chunk.",
             "rag.py", lines("rag.py", 267, 283),
             "29_rag_top_k_sweep.png", 29, "Top-k sweep showing k=5 as the first complete and correct setting.")

    doc.add_page_break()
    doc.add_heading("4.7 RAG comparison analysis", level=2)
    analysis = (HW4 / "RAG_ANALYSIS.md").read_text(encoding="utf-8").split("\n\n", 1)[1]
    for paragraph in analysis.split("\n\n"):
        doc.add_paragraph(paragraph.strip())
    doc.add_paragraph("Word count: 361")

    doc.add_page_break()
    doc.add_heading("Reproducibility and self-check", level=1)
    doc.add_paragraph(
        "The repository contains the schema, deterministic seeder, React and FastAPI code, Postman "
        "collection, 180-request raw N+1 CSV and JSON files, EXPLAIN outputs, RAG results, retrieval "
        "printouts, evaluation data, AI_USE.md, METRICS.md, RUN_LOG.txt, and verification.json."
    )
    add_code(doc, "Smoke-test entry point", lines("scripts/verify_hw04.py", 49, 69) + "\n...\n" + lines("scripts/verify_hw04.py", 118, 145))
    verification = json.loads((HW4 / "verification.json").read_text(encoding="utf-8"))
    rows = [[item["name"], "PASS" if item["passed"] else "FAIL", item["details"]] for item in verification["checks"]]
    add_table(doc, ["Check", "Result", "Details"], rows, widths=[2.3, 0.8, 3.4])
    doc.add_paragraph(
        "Result: 11 of 11 objective checks passed. verification.json records code-and-results commit "
        f"{commit_hash}, so the smoke-test evidence identifies the submitted implementation."
    )

    doc.add_page_break()
    doc.add_heading("AI-use disclosure", level=1)
    ai_text = (HW4 / "AI_USE.md").read_text(encoding="utf-8")
    for block in ai_text.split("\n\n"):
        stripped = block.strip()
        if not stripped or stripped.startswith("# AI-Use"):
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        else:
            doc.add_paragraph(stripped)

    core = doc.core_properties
    core.title = "DATA 260 Homework 4"
    core.subject = "React, MySQL, N+1 query tuning, and grounded RAG"
    core.author = "Vedant Vartak"
    core.keywords = "DATA 260, Homework 4, React, FastAPI, MySQL, RAG"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--commit-hash", default="PENDING FINAL CODE AND RESULTS COMMIT")
    args = parser.parse_args()
    build(args.commit_hash)
